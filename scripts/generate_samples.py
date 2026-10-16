"""Generates the sample/test corpus (needs requirements-dev.txt: reportlab, matplotlib).

    python scripts/generate_samples.py [output_dir]      # default: tests/samples
"""
from __future__ import annotations

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pymupdf
from docx import Document
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FONT = font_manager.findfont("DejaVu Sans")


def _font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def _table(data, widths=None):
    t = Table(data, colWidths=widths)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.8, colors.black),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                           ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    return t


def revenue_chart(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(6, 4), dpi=150)
    regions, share = ["North America", "Europe", "Asia"], [42, 31, 27]
    bars = ax.bar(regions, share, color=["#1f77b4", "#ff7f0e", "#2ca02c"])
    for b, v in zip(bars, share):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.8, f"{v}%", ha="center", fontsize=12)
    ax.set_title("Revenue Share by Region 2025")
    ax.set_ylabel("Share of revenue (%)")
    ax.set_ylim(0, 50)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def company_pdf(path: Path, company: str, quarters, objectives: list[str], chart: Path | None, outlook: str) -> None:
    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    story = [Paragraph(f"{company} Annual Report 2025", st["Title"]), Spacer(1, 12),
             Paragraph("Introduction", st["Heading1"]),
             Paragraph(f"{company} is a mid-sized technology firm. This report summarises financial performance, "
                       "project objectives and the outlook for the coming year.", st["BodyText"]),
             Paragraph("Project Objectives", st["Heading1"])]
    story += [Paragraph(f"{i + 1}. {o}", st["BodyText"]) for i, o in enumerate(objectives)]
    story += [PageBreak(), Paragraph("Quarterly Financials", st["Heading1"]),
              Paragraph("The table below lists revenue, expenses and net profit per quarter (all values in USD).", st["BodyText"]),
              Spacer(1, 8), _table(quarters), PageBreak()]
    if chart:
        story += [Paragraph("Regional Performance", st["Heading1"]),
                  Paragraph("Figure 1 shows how revenue was distributed geographically.", st["BodyText"]),
                  RLImage(str(chart), width=380, height=253), PageBreak()]
    story += [Paragraph("Risks and Outlook", st["Heading1"]), Paragraph(outlook, st["BodyText"])]
    doc.build(story)


def employees_docx(path: Path, diagram: Path) -> None:
    d = Document()
    d.add_heading("Employee Directory", level=1)
    d.add_paragraph("This directory lists staff, their departments and annual salaries in USD.")
    rows = [("Name", "Department", "Salary (USD)", "Years"), ("Alice Johnson", "Engineering", "120,000", "6"),
            ("Bob Smith", "Marketing", "85,000", "3"), ("Carol Diaz", "Engineering", "135,000", "9"),
            ("David Lee", "Sales", "70,000", "2"), ("Emma Wilson", "HR", "78,000", "5")]
    t = d.add_table(rows=len(rows), cols=4)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v
    d.add_heading("Organisation Chart", level=1)
    d.add_paragraph("The chart below shows the reporting structure.")
    d.add_picture(str(diagram))
    d.save(path)


def diagram_png(path: Path) -> None:
    img = Image.new("RGB", (900, 260), "white")
    dr = ImageDraw.Draw(img)
    labels = ["CEO", "CTO", "Engineering"]
    for i, lab in enumerate(labels):
        x = 40 + i * 290
        dr.rectangle([x, 80, x + 220, 170], outline="black", width=4)
        dr.text((x + 110, 125), lab, fill="black", font=_font(30), anchor="mm")
        if i < 2:
            dr.line([x + 220, 125, x + 290, 125], fill="black", width=4)
            dr.polygon([(x + 290, 125), (x + 270, 113), (x + 270, 137)], fill="black")
    img.save(path)


def text_image(path: Path, lines: list[str], size=(1100, 520), rotate: float = 0.0, font_size=40) -> None:
    img = Image.new("RGB", size, "white")
    dr = ImageDraw.Draw(img)
    y = 40
    for ln in lines:
        dr.text((50, y), ln, fill="black", font=_font(font_size))
        y += font_size + 22
    if rotate:
        img = img.rotate(rotate, expand=False, fillcolor="white", resample=Image.BICUBIC)
    img.save(path)


def scanned_pdf(path: Path, png: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(30, 60, 565, 340), filename=str(png))  # image only: no text layer
    doc.save(path)
    doc.close()


def generate_all(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    chart = out / "_chart.png"
    revenue_chart(chart)

    files["acme"] = out / "acme_report_2025.pdf"
    company_pdf(files["acme"], "Acme Corp",
                [["Quarter", "Revenue (USD)", "Expenses (USD)", "Net Profit (USD)"],
                 ["Q1", "1,200,000", "900,000", "300,000"], ["Q2", "1,350,000", "1,000,000", "350,000"],
                 ["Q3", "1,100,000", "1,250,000", "(150,000)"], ["Q4", "1,600,000", "1,100,000", "500,000"],
                 ["Total", "5,250,000", "4,250,000", "1,000,000"]],
                ["Expand the cloud platform to three new markets.", "Reduce customer churn below 5 percent.",
                 "Launch the mobile application by the third quarter."],
                chart, "The main risk is currency volatility. Management expects moderate growth next year.")

    files["beta"] = out / "beta_ltd_report.pdf"
    company_pdf(files["beta"], "Beta Ltd",
                [["Quarter", "Revenue (USD)", "Expenses (USD)", "Net Profit (USD)"],
                 ["Q1", "800,000", "600,000", "200,000"], ["Q2", "900,000", "650,000", "250,000"],
                 ["Q3", "950,000", "700,000", "250,000"], ["Q4", "1,000,000", "720,000", "280,000"]],
                ["Open a regional office in Singapore.", "Train 200 engineers in machine learning."],
                None, "Beta Ltd anticipates supply chain delays and expects steady growth.")

    diagram = out / "_orgchart.png"
    diagram_png(diagram)
    files["docx"] = out / "employees.docx"
    employees_docx(files["docx"], diagram)
    files["diagram"] = out / "flow_diagram.png"
    diagram_png(files["diagram"])

    memo = out / "_memo.png"
    text_image(memo, ["MEMORANDUM", "Date: 12 March 2025", "The board approved a budget of $75,000",
                      "for the Kerala pilot project.", "Project lead: Anita Menon"], size=(1200, 560), rotate=1.5)
    files["scanned"] = out / "scanned_memo.pdf"
    scanned_pdf(files["scanned"], memo)

    files["invoice"] = out / "invoice.png"
    text_image(files["invoice"], ["INVOICE INV-2045", "Customer: Northwind Traders", "Total Due: $1,284.50",
                                  "Due Date: 30 April 2025"], size=(1000, 460))
    files["txt"] = out / "notes.txt"
    files["txt"].write_text("Meeting notes\n\nThe team agreed to migrate the database to PostgreSQL in June. "
                            "Marta will lead the migration and Raj will handle testing.\n", encoding="utf-8")
    return files


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/samples")
    for k, v in generate_all(target).items():
        print(f"{k:9s} -> {v}")
